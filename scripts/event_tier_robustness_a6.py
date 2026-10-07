"""A6 event-tier robustness re-cut of the locked A5 ledger (prereg
docs/PREREG_ADDENDUM_A6_EVENT_TIER_ROBUSTNESS_2026-10-07.md, lock 8769356e).
Mutation content only; no MICs. Genes re-read from pinned sources only."""
import json, re, zipfile, io
from collections import defaultdict
from pathlib import Path
import openpyxl

BASE = Path(__file__).resolve().parents[1]
ADOPTED = {'BAC5','CAP18','CP1','HBD3','IND','LL37','PEX','PGLA','PLEU','PR39','R8','TPII'}
TIER_A_EXACT = {'SNP','SNPs','single nucleotide polymorphism','insertion',
                'insertion (IC3-like element, 1.4 kb)','intergenic mutation','short_deletion'}
TIER_A_RE = re.compile(r'^(Del|Ins) \d+ nt$')
TIER_B_EXACT = {'large_deletion'}

def tier(mutation_type):
    if mutation_type in TIER_B_EXACT: return 'B'
    if mutation_type in TIER_A_EXACT or TIER_A_RE.match(mutation_type): return 'A'
    raise ValueError(f'unrecognized mutation_type {mutation_type!r}; halt per A6.1')

def clean_gene(raw):
    g = re.sub(r'[\xa0\s]', ' ', str(raw)).strip()
    return re.sub(r'\s*[→←]+\s*', '', g).strip()

def spohn_sheets():
    z = zipfile.ZipFile(BASE/'data/raw/spohn2019_supp.zip')
    return openpyxl.load_workbook(io.BytesIO(z.read('41467_2019_12364_MOESM9_ESM.xlsx')),
                                  read_only=True, data_only=True)

def spohn_blocks():
    """Unique large-deletion blocks keyed by (Length bp, Start bp); one structural event each."""
    s = next(s for s in spohn_sheets() if s.title == 'large_deletion')
    rows = list(s.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and any(str(c).strip() == 'Strain' for c in r if c))
    hdr = [str(c).strip() if c else '' for c in rows[hi]]
    si, li, bi, gi = hdr.index('Strain'), hdr.index('Length (bp)'), hdr.index('Start bp'), hdr.index('List of deleted genes')
    blocks = {}
    for r in rows[hi+1:]:
        if not r or not r[si]: continue
        strain = str(r[si]).strip()
        amp = strain.rsplit('_', 1)[0]
        if amp not in ADOPTED: continue
        key = (str(r[li]).strip(), str(r[bi]).strip())
        b = blocks.setdefault(key, {'lines': set(), 'amps': set(), 'genes': set()})
        b['lines'].add(strain); b['amps'].add(amp)
        for g in re.split(r'[,/]', str(r[gi])):
            g = clean_gene(g)
            if g and not g.startswith('['):
                b['genes'].add(g)
    return blocks

def spohn_direct_starts():
    """(sheet, line, gene) -> Start bp for direct-mutation identity checks (A6.3)."""
    out = {}
    for sheet, gcol in (('SNPs','Gene'), ('short_deletion','Gene'), ('insertion','Gene'), ('intergenic mutation','Gene')):
        s = next(s for s in spohn_sheets() if s.title == sheet)
        rows = list(s.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and any(str(c).strip() == 'Strain' for c in r if c))
        hdr = [str(c).strip() if c else '' for c in rows[hi]]
        si, gi, bi = hdr.index('Strain'), hdr.index(gcol), hdr.index('Start bp')
        for r in rows[hi+1:]:
            if not r or not r[si]: continue
            strain = str(r[si]).strip()
            if strain.rsplit('_',1)[0] not in ADOPTED: continue
            for g in re.split(r'[,/]', str(r[gi])):
                g = clean_gene(g)
                if g and not g.startswith('['):
                    out[(sheet, strain, g)] = str(r[bi]).strip()
    return out

def run():
    ledger = json.load(open(BASE/'results/mutation_convergence_a5.json'))
    recs = ledger['records']
    evo = [r for r in recs if r.get('scope','evolved') == 'evolved']
    # Guard: recompute A5 totals from the ledger itself (A6 verification)
    counts = {'total_records': len(recs),
              'spohn_evolved': sum(1 for r in evo if r['study']=='spohn2019'),
              'blanco_evolved': sum(1 for r in evo if r['study']=='blanco2020'),
              'blanco_control_excluded': sum(1 for r in recs if r.get('scope')=='control'),
              'bac7_evolved': sum(1 for r in evo if r['study']=='bac7_2023')}
    assert counts == ledger['counts'], ('A5 guard mismatch', counts, ledger['counts'])
    # Guard: recompute A5 within-study recurrence (A5 logic, first method)
    groups, types = defaultdict(set), defaultdict(set)
    for r in evo:
        groups[(r['study'],r['amp'],r['gene'])].add(r['line'])
        types[(r['study'],r['amp'],r['gene'])].add(r['mutation_type'])
    within = sorted(({'study':k[0],'amp':k[1],'gene':k[2],'n_lines':len(v),
                      'mutation_types':sorted(types[k]),
                      'driven_by_large_deletion':types[k]=={'large_deletion'}}
                     for k,v in groups.items() if len(v)>=2),
                    key=lambda x:(-x['n_lines'],x['gene']))
    assert within == ledger['within_study_recurrence'], 'A5 within-study guard mismatch'
    # Tier assignment (halts on any unrecognized type string)
    for r in recs: r['tier'] = tier(r['mutation_type'])
    # Strict Tier A cross-study recurrence (A6.2)
    strict = defaultdict(set)
    for r in evo:
        if r['tier'] == 'A': strict[r['gene'].lower()].add(r['study'])
    strict_cross = sorted(g for g, s in strict.items() if len(s) >= 2)
    # Event model (A6.3): blocks coalesced by (Length, Start); direct hits keyed by
    # (gene, type, start) for Spohn (Start bp from pinned sheet), per (line,gene,type) elsewhere.
    blocks = spohn_blocks()
    starts = spohn_direct_starts()
    events = defaultdict(set)           # gene.lower() -> {(tier, event_id)}
    event_members = defaultdict(set)    # event_id -> lines
    sbmA_detail = []
    for r in evo:
        g = r['gene'].lower()
        if r['study'] == 'spohn2019' and r['tier'] == 'B':
            continue  # handled via blocks
        if r['study'] == 'spohn2019':
            st = starts.get((r['mutation_type'], r['line'], r['gene']), 'nostart')
            eid = f"A|spohn2019|{r['mutation_type']}|{r['gene']}|{st}"
            indep = 'independent unless identical coordinates; start recorded'
        else:
            eid = f"A|{r['study']}|{r['line']}|{r['gene']}|{r['mutation_type']}"
            indep = 'one event per independent strain/clone; coordinate check not available in A5 extraction (disclosed)'
        events[g].add(('A', eid)); event_members[eid].add(r['line'])
        if g == 'sbma':
            sbmA_detail.append({'study': r['study'], 'amp': r['amp'], 'line': r['line'],
                                'mutation_type': r['mutation_type'], 'tier': 'A',
                                'event_id': eid, 'independence': indep})
    for (ln, st), b in sorted(blocks.items()):
        eid = f"B|spohn2019|block|L{ln}|S{st}"
        for g in b['genes']:
            events[g.lower()].add(('B', eid))
            if g.lower() == 'sbma':
                sbmA_detail.append({'study': 'spohn2019', 'amp': '+'.join(sorted(b['amps'])),
                                    'line': '+'.join(sorted(b['lines'])),
                                    'mutation_type': f'large_deletion {ln} bp @ {st}',
                                    'tier': 'B', 'event_id': eid,
                                    'independence': 'coalesced structural event per A6.3'})
        event_members[eid] |= b['lines']
    # All-recurrent-genes comparison (A6.4), descriptive only
    table = []
    for w in within:
        g = w['gene'].lower()
        ev = events.get(g, set())
        studies = sorted({r['study'] for r in evo if r['gene'].lower() == g})
        table.append({'gene': w['gene'], 'study': w['study'], 'amp': w['amp'],
                      'n_lines': w['n_lines'], 'mutation_types': w['mutation_types'],
                      'driven_by_large_deletion': w['driven_by_large_deletion'],
                      'tierA_lines': len({r['line'] for r in evo if r['gene'].lower()==g and r['tier']=='A'}),
                      'tierB_lines': len({r['line'] for r in evo if r['gene'].lower()==g and r['tier']=='B'}),
                      'direct_events': sum(1 for t,_ in ev if t=='A'),
                      'block_events': sum(1 for t,_ in ev if t=='B'),
                      'n_studies': len(studies), 'studies': studies,
                      'strict_cross_study': g in strict_cross})
    table.sort(key=lambda x: (-x['n_studies'], -x['n_lines'], x['gene']))
    # Second pass, different method: Counter-based strict recount
    from collections import Counter
    c = Counter()
    for r in evo:
        if r['tier'] == 'A': c[(r['gene'].lower(), r['study'])] += 1
    strict2 = sorted(g for (g, _), n in c.items() if g in {gg for (gg, ss) in c} )
    genes2 = defaultdict(set)
    for (g, s) in c: genes2[g].add(s)
    strict2 = sorted(g for g, ss in genes2.items() if len(ss) >= 2)
    assert strict2 == strict_cross, 'second-pass strict mismatch'
    out = {'prereg': 'A6', 'lock_commit': '8769356e6204dec1291720f5fa843a0755a2f01e',
           'guard_a5_counts': counts, 'guard_within_study_match': True,
           'tier_map': {t: tier(t) for t in sorted({r['mutation_type'] for r in recs})},
           'strict_cross_study_tierA': strict_cross,
           'a5_all_tier_cross_study': [m['gene'] for m in ledger['cross_study_exact_matches']],
           'sbmA_event_ledger': sorted(sbmA_detail, key=lambda d: (d['tier'], d['event_id'])),
           'spohn_block_count': len(blocks),
           'all_recurrent_genes': table}
    (BASE/'results/event_tier_robustness_a6.json').write_text(json.dumps(out, indent=2))
    print(json.dumps({k: out[k] for k in ('guard_a5_counts','strict_cross_study_tierA',
                      'a5_all_tier_cross_study','spohn_block_count')}, indent=2))
    print('sbmA events:')
    for d in out['sbmA_event_ledger']: print(' ', d['tier'], d['event_id'], '|', d['amp'], d['line'])
    top = [t for t in table if t['gene'].lower()=='sbma']
    print('sbmA rows in all-recurrent table:', json.dumps(top, indent=2))

if __name__ == '__main__':
    run()
